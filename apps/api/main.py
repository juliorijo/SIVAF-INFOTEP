from __future__ import annotations

from datetime import datetime, timezone
import hashlib
import csv
import io
import json
import os
from pathlib import Path
import re
import secrets
import time
import unicodedata
from zipfile import BadZipFile
from typing import Any, Dict, List
from uuid import UUID, uuid4

from fastapi import BackgroundTasks, Depends, FastAPI, File, Form, HTTPException, Request, Response, UploadFile, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse
import pypdfium2 as pdfium
from pydantic import BaseModel, Field
from sqlalchemy import delete, func, select
from starlette.concurrency import run_in_threadpool

from apps.api.auth import verify_password
from apps.api.database import (
    AuditEventORM,
    FormativeActionORM,
    FormativeActionImportORM,
    HumanReviewORM,
    ProcessingJobORM,
    SessionLocal,
    UserORM,
    UserSessionORM,
    ValidationResultORM,
    init_db,
)
from packages.document_processing.pdf_service import (
    MAX_PDF_RASTER_PIXELS,
    PdfIngestionService,
    PdfUploadLimitError,
)
from packages.domain.enums import JobStatus, ResultOutcome, RuleCode
from packages.ocr.ai_assisted_analyzer import AIAssistedOCRAnalyzer, create_ai_assisted_analyzer
from packages.ocr.registry import load_ocr_provider
from packages.validation.dominican_cedula import (
    is_dominican_cedula_checksum_valid,
    normalize_dominican_cedula_candidate,
)

app = FastAPI(
    title="SIVAF | Sistema de Verificación de Acciones Formativas",
    version="0.1.0",
    description="API de SIVAF para la verificación documental de acciones formativas.",
)
app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:3000",
        "http://127.0.0.1:3000",
        "http://localhost:5500",
        "http://127.0.0.1:5500",
    ],
    allow_credentials=True,
    allow_methods=["GET", "POST", "PUT", "DELETE", "OPTIONS"],
    allow_headers=["Content-Type"],
)

pdf_service = PdfIngestionService()
ocr_provider = None
ocr_provider_diagnostic = ""
ai_ocr_analyzer: AIAssistedOCRAnalyzer | None = None
ai_ocr_diagnostic = ""
AUTH_COOKIE_NAME = "sivaf_session"
AUTH_SESSION_SECONDS = 8 * 60 * 60
AUTH_COOKIE_SECURE = os.getenv("SIVAF_COOKIE_SECURE", "false").lower() == "true"
MAX_ACTION_IMPORT_BYTES = 20 * 1024 * 1024
MAX_UPLOAD_METADATA_BYTES = 16 * 1024
MAX_ACTION_IMPORT_ROWS = 50_000
ACTION_FIELD_ALIASES = {
    "action_code": {
        "codigo", "codigo accion", "codigo de accion", "codigo accion formativa",
        "numero accion", "numero de accion", "no accion", "nro accion",
        "num accion", "id accion", "codigo c",
    },
    "name": {
        "accion formativa", "nombre accion", "nombre de accion", "nombre",
        "nombre actividad", "nombre del curso", "curso", "denominacion",
        "denominacion accion", "accion f",
    },
    "start_date": {"fecha inicio", "fecha de inicio", "inicio", "fecha inicial", "fecha ini"},
    "status": {"estado", "estatus", "situacion", "status"},
}


def load_local_env_file() -> None:
    for filename in (".env", ".env.local"):
        env_path = Path(__file__).resolve().parents[2] / filename
        if not env_path.is_file():
            continue
        for raw_line in env_path.read_text(encoding="utf-8").splitlines():
            line = raw_line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            key = key.strip()
            value = value.strip().strip('"').strip("'")
            if key:
                os.environ[key] = value


@app.on_event("startup")
def startup_db() -> None:
    global ai_ocr_analyzer, ai_ocr_diagnostic, ocr_provider, ocr_provider_diagnostic
    load_local_env_file()
    init_db()
    try:
        ocr_provider = load_ocr_provider()
        if ocr_provider is None:
            ocr_provider_diagnostic = (
                "No se encontró un motor OCR local disponible. "
                "Instala Tesseract OCR y define TESSERACT_CMD o configura OCR_PROVIDER_FACTORY."
            )
        else:
            ocr_provider_diagnostic = ""
    except Exception as exc:
        ocr_provider = None
        ocr_provider_diagnostic = f"Error de configuración OCR: {exc}"
    try:
        ai_ocr_analyzer = create_ai_assisted_analyzer()
        if ai_ocr_analyzer.is_configured():
            ai_ocr_diagnostic = ""
        else:
            ai_ocr_diagnostic = (
                "AI OCR deshabilitado: define AI_OCR_BASE_URL y AI_OCR_MODEL para activar la revisión asistida."
            )
    except Exception as exc:
        ai_ocr_analyzer = None
        ai_ocr_diagnostic = f"Error de configuración AI OCR: {exc}"
    
    # Inicializar permisos y roles por defecto (Fase 1)
    from apps.api.permissions import create_default_permissions, create_default_roles
    session = SessionLocal()
    try:
        create_default_permissions(session)
        create_default_roles(session)
    finally:
        session.close()


init_db()

# ============================================================================
# INCLUIR ROUTER DE FASE 2 (Permisos & Multi-usuario)
# ============================================================================
from apps.api.routes_phase2 import router as phase2_router
app.include_router(phase2_router)


class CreateJobRequest(BaseModel):
    source_pdf_path: str = Field(..., min_length=1, description="Ruta del PDF a procesar")
    metadata: Dict[str, Any] = Field(default_factory=dict, description="Metadatos adicionales del trabajo")


class ProcessJobRequest(BaseModel):
    force: bool = Field(default=False, description="Forzar el re-procesamiento del PDF")
    metadata: Dict[str, Any] = Field(default_factory=dict, description="Metadatos del procesamiento")


class HumanReviewRequest(BaseModel):
    comments: str = Field(default="", max_length=4000)
    final_status: str | None = Field(default=None)


class LoginRequest(BaseModel):
    username: str = Field(..., min_length=1, max_length=80)
    password: str = Field(..., min_length=1, max_length=1024)


FINAL_STATUSES = {"APROBADA", "DEVUELTA", "MODIFICACION", "OTRO"}


def get_current_user(request: Request) -> UserORM:
    token = request.cookies.get(AUTH_COOKIE_NAME)
    if not token:
        raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Authentication required")

    token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
    with SessionLocal() as session:
        auth_session = session.get(UserSessionORM, token_hash)
        if auth_session is None:
            raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Session expired or invalid")
        if auth_session.expires_at <= time.time():
            session.delete(auth_session)
            session.commit()
            raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Session expired or invalid")

        user = session.get(UserORM, auth_session.user_id)
        if user is None or not user.is_active:
            session.delete(auth_session)
            session.commit()
            raise HTTPException(status_code=status.HTTP_401_UNAUTHORIZED, detail="Session expired or invalid")
        return user


def require_admin(current_user: UserORM = Depends(get_current_user)) -> UserORM:
    if current_user.role != "admin":
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Administrator role required")
    return current_user


def get_accessible_job(session, job_id: str, current_user: UserORM) -> ProcessingJobORM:
    job = session.get(ProcessingJobORM, job_id)
    if job is None or (
        current_user.role != "admin"
        and job.created_by != current_user.id
    ):
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Processing job not found")
    return job


def audit_evidence_access(job_id: str, actor: str, action: str, details: Dict[str, Any]) -> None:
    with SessionLocal() as session:
        session.add(AuditEventORM(
            id=str(uuid4()),
            job_id=job_id,
            actor=actor,
            action=action,
            details=sanitize_json(details),
            created_at=datetime.now(timezone.utc),
        ))
        session.commit()


def normalize_action_header(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(character for character in text if not unicodedata.combining(character))
    return " ".join(re.sub(r"[^a-zA-Z0-9]+", " ", text.lower()).split())


def normalize_spreadsheet_value(value: Any) -> Any:
    if hasattr(value, "isoformat"):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def parse_action_import(filename: str, content: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        try:
            rows = list(csv.reader(io.StringIO(content.decode("utf-8-sig"), newline="")))
        except UnicodeDecodeError as exc:
            raise HTTPException(status_code=400, detail="El CSV debe estar guardado como UTF-8.") from exc
    elif suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
            from openpyxl.utils.exceptions import InvalidFileException

            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            worksheet = workbook.active
            rows = list(worksheet.iter_rows(values_only=True))
            workbook.close()
        except ImportError as exc:
            raise HTTPException(status_code=500, detail="Falta instalar la dependencia openpyxl para leer Excel.") from exc
        except (BadZipFile, InvalidFileException, OSError, ValueError) as exc:
            raise HTTPException(status_code=400, detail="No se pudo leer el archivo Excel .xlsx.") from exc
    else:
        raise HTTPException(status_code=400, detail="Usa un archivo .xlsx o .csv exportado desde Drive.")

    if not rows:
        raise HTTPException(status_code=400, detail="El archivo no contiene filas.")

    def locate_header(values: Any) -> tuple[list[str], dict[str, int], dict[str, int]] | None:
        candidate_headers = [str(value).strip().strip('"') if value is not None else "" for value in values]
        header_indexes = {
            normalize_action_header(header): index
            for index, header in enumerate(candidate_headers)
            if normalize_action_header(header)
        }
        mapped = {}
        for field, aliases in ACTION_FIELD_ALIASES.items():
            exact = next((header_indexes[alias] for alias in aliases if alias in header_indexes), None)
            if exact is None:
                partial_matches = [
                    index for normalized, index in header_indexes.items()
                    if len(normalized) >= 5 and any(alias.startswith(normalized) or normalized.startswith(alias) for alias in aliases)
                ]
                exact = partial_matches[0] if len(set(partial_matches)) == 1 else None
            mapped[field] = exact
        if mapped["action_code"] is not None:
            return candidate_headers, header_indexes, mapped
        return None

    header_location = None
    for row_index, values in enumerate(rows[:50]):
        header_location = locate_header(values)
        if header_location:
            break
    if header_location is None:
        raise HTTPException(
            status_code=400,
            detail="No se encontró la fila de encabezados con un código de acción. Revisa que el CSV incluya sus columnas.",
        )
    headers, _, mapped_indexes = header_location
    if mapped_indexes["action_code"] is None:
        raise HTTPException(
            status_code=400,
            detail="No se encontró una columna de código de acción (por ejemplo: Código de acción o No. Acción).",
        )

    records: list[dict[str, Any]] = []
    seen_codes: set[str] = set()
    for row_number, values in enumerate(rows[row_index + 1:], start=row_index + 2):
        if not any(value is not None and str(value).strip() for value in values):
            continue
        if len(records) >= MAX_ACTION_IMPORT_ROWS:
            raise HTTPException(status_code=400, detail=f"El archivo supera el máximo de {MAX_ACTION_IMPORT_ROWS} registros.")

        fields = {
            header: normalize_spreadsheet_value(values[index] if index < len(values) else None)
            for index, header in enumerate(headers)
            if header
        }
        code_index = mapped_indexes["action_code"]
        action_code = str(values[code_index] if code_index < len(values) else "").strip()
        if not action_code:
            raise HTTPException(status_code=400, detail=f"Falta el código de acción en la fila {row_number}.")
        if action_code in seen_codes:
            raise HTTPException(status_code=400, detail=f"El código de acción {action_code} está duplicado en el archivo.")
        seen_codes.add(action_code)

        def mapped_value(field: str) -> str | None:
            index = mapped_indexes[field]
            if index is None or index >= len(values) or values[index] is None:
                return None
            return str(normalize_spreadsheet_value(values[index])).strip() or None

        records.append({
            "action_code": action_code,
            "name": mapped_value("name") or "",
            "start_date": mapped_value("start_date"),
            "status": mapped_value("status"),
            "fields": fields,
        })
    if not records:
        raise HTTPException(status_code=400, detail="No se encontraron registros de acciones formativas.")
    columns = [header for header in headers if header]
    return records, columns


def is_managed_evidence_available(source_pdf_path: str) -> bool:
    evidence_path = Path(source_pdf_path).resolve()
    try:
        evidence_path.relative_to(pdf_service.storage_dir.resolve())
    except ValueError:
        return False
    return evidence_path.is_file()


def render_pdf_page_preview(source_pdf_path: str, page_number: int) -> tuple[bytes, int]:
    with pdfium.PdfDocument(source_pdf_path) as document:
        page_count = len(document)
        if page_number < 1 or page_number > page_count:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"PDF page must be between 1 and {page_count}.",
            )
        page = document[page_number - 1]
        scale = 1.5
        if page.get_width() * page.get_height() * scale**2 > MAX_PDF_RASTER_PIXELS:
            page.close()
            raise HTTPException(
                status_code=status.HTTP_413_CONTENT_TOO_LARGE,
                detail="PDF page exceeds the safe preview raster limit.",
            )
        bitmap = page.render(scale=scale)
        try:
            with bitmap.to_pil() as image:
                output = io.BytesIO()
                image.save(output, format="PNG", optimize=True)
                return output.getvalue(), page_count
        finally:
            bitmap.close()
            page.close()


def serialize_processing_job(row: ProcessingJobORM, review: HumanReviewORM | None = None) -> Dict[str, Any]:
    return {
        "id": row.id,
        "source_pdf_name": Path(row.source_pdf_path).name,
        "original_file_hash": row.original_file_hash,
        "status": row.status,
        "queued_at": row.queued_at.isoformat() if row.queued_at else None,
        "started_at": row.started_at.isoformat() if row.started_at else None,
        "finished_at": row.finished_at.isoformat() if row.finished_at else None,
        "created_by": row.created_by,
        "metadata": row.payload_metadata or {},
        "evidence_available": is_managed_evidence_available(row.source_pdf_path),
        "human_review": {
            "reviewer": review.reviewer,
            "comments": review.comments,
            "final_status": review.final_status,
            "reviewed_at": review.reviewed_at.isoformat(),
        } if review else None,
    }


def serialize_validation_result(row: ValidationResultORM) -> Dict[str, Any]:
    evidence_ref = row.evidence_ref
    if evidence_ref and re.match(r"^(?:[A-Za-z]:[\\/]|/)", evidence_ref):
        evidence_ref = Path(evidence_ref).name
    return {
        "id": row.id,
        "job_id": row.job_id,
        "participant_id": row.participant_id,
        "identity_document_id": row.identity_document_id,
        "rule_id": row.rule_id,
        "outcome": row.outcome,
        "reason_code": row.reason_code,
        "message": row.message,
        "evidence_ref": evidence_ref,
        "created_at": row.created_at.isoformat() if row.created_at else None,
    }


def find_cedula_candidates(text: str) -> list[str]:
    if not text:
        return []
    pattern = re.compile(
        r"(?<![A-Z0-9])(?:[0-9OQDILSBGZ]{3}[\s\-.:_/,;]*[0-9OQDILSBGZ]{7}[\s\-.:_/,;]*[0-9OQDILSBGZ]|[0-9OQDILSBGZ]{11})(?![A-Z0-9])",
        re.IGNORECASE,
    )
    candidates = {
        normalized
        for raw_match in pattern.findall(text.upper())
        if (normalized := normalize_dominican_cedula_candidate(raw_match))
    }
    return sorted(candidates)


def mask_cedula_candidate(number: str) -> str:
    compact = re.sub(r"\D", "", number or "")
    if len(compact) != 11:
        return "***"
    return f"***-*******-{compact[-1]} · ****{compact[-4:]}"


def summarize_document_candidates(pages) -> Dict[str, Any]:
    roster_by_page: dict[int, set[str]] = {}
    evidence_by_page: dict[int, set[str]] = {}
    low_confidence_evidence_candidates: set[str] = set()
    confident_roster_numbers: set[str] = set()
    confident_evidence_numbers: set[str] = set()
    roster_checksum_valid_numbers: set[str] = set()
    evidence_checksum_valid_numbers: set[str] = set()
    roster_checksum_invalid_numbers: set[str] = set()
    evidence_checksum_invalid_numbers: set[str] = set()
    candidate_profiles: dict[str, Dict[str, Any]] = {}

    for page in pages:
        numbers = set(find_cedula_candidates(page.text))
        if not numbers:
            continue

        text = page.text.upper()
        has_document_label = bool(re.search(r"C[EÉ]DULA|DOCUMENTO|IDENTIDAD", text))
        has_list_label = bool(re.search(r"LISTA|N[ÓO]MINA|RELACI[ÓO]N", text))
        has_participant_label = bool(re.search(r"PARTICIPANTE|ALUMNO|ESTUDIANTE", text))
        has_name_label = bool(re.search(r"NOMBRE|APELLIDO", text))
        is_roster_page = has_document_label and (
            (has_list_label and has_participant_label)
            or (has_participant_label and has_name_label and len(numbers) > 1)
        )

        if is_roster_page:
            roster_by_page[page.page_number] = numbers
            for number in numbers:
                profile = candidate_profiles.setdefault(number, {
                    "roster_pages": set(),
                    "document_pages": set(),
                    "roster_confident": False,
                    "document_confident": False,
                    "document_low_confidence": False,
                })
                profile["roster_pages"].add(page.page_number)
                if is_dominican_cedula_checksum_valid(number):
                    roster_checksum_valid_numbers.add(number)
                    if page.confidence is not None and page.confidence >= 0.8:
                        confident_roster_numbers.add(number)
                        profile["roster_confident"] = True
                else:
                    roster_checksum_invalid_numbers.add(number)
        else:
            evidence_by_page[page.page_number] = numbers
            for number in numbers:
                profile = candidate_profiles.setdefault(number, {
                    "roster_pages": set(),
                    "document_pages": set(),
                    "roster_confident": False,
                    "document_confident": False,
                    "document_low_confidence": False,
                })
                profile["document_pages"].add(page.page_number)
                if is_dominican_cedula_checksum_valid(number):
                    evidence_checksum_valid_numbers.add(number)
                    if page.confidence is not None and page.confidence >= 0.8:
                        confident_evidence_numbers.add(number)
                        profile["document_confident"] = True
                else:
                    evidence_checksum_invalid_numbers.add(number)
            if page.confidence is None or page.confidence < 0.8:
                low_confidence_evidence_candidates.update(numbers)
                for number in numbers:
                    profile = candidate_profiles.setdefault(number, {
                        "roster_pages": set(),
                        "document_pages": set(),
                        "roster_confident": False,
                        "document_confident": False,
                        "document_low_confidence": False,
                    })
                    profile["document_low_confidence"] = True

    roster_numbers = set().union(*roster_by_page.values()) if roster_by_page else set()
    evidence_numbers = set().union(*evidence_by_page.values()) if evidence_by_page else set()
    confident_matches = confident_roster_numbers & confident_evidence_numbers
    all_candidates = sorted(roster_numbers | evidence_numbers)
    participant_review_items: list[Dict[str, Any]] = []
    for index, number in enumerate(all_candidates, start=1):
        profile = candidate_profiles.get(number, {})
        roster_pages = sorted(profile.get("roster_pages", set()))
        document_pages = sorted(profile.get("document_pages", set()))
        roster_confident = bool(profile.get("roster_confident"))
        document_confident = bool(profile.get("document_confident"))
        checksum_valid = is_dominican_cedula_checksum_valid(number)
        if not checksum_valid:
            review_status = "REQUIERE_REVISION"
            review_reason = "Dígito de control no coincide."
        elif roster_confident and document_confident:
            review_status = "OK_AUTOMATICO"
            review_reason = "Coincidencia OCR confiable entre lista y documento."
        elif roster_confident and not document_confident:
            review_status = "REQUIERE_REVISION"
            review_reason = "No se encontró coincidencia confiable en páginas de documento."
        elif not roster_confident and document_confident:
            review_status = "REQUIERE_REVISION"
            review_reason = "Detectado en documento, pero no en lista con confianza alta."
        else:
            review_status = "REQUIERE_REVISION"
            review_reason = "Lectura OCR de baja confianza; validar manualmente."
        participant_review_items.append({
            "reference": f"P-{index:03d}",
            "cedula_masked": mask_cedula_candidate(number),
            "review_status": review_status,
            "review_reason": review_reason,
            "roster_pages": roster_pages,
            "document_pages": document_pages,
            "roster_sheets": sorted({(page + 1) // 2 for page in roster_pages}),
            "document_sheets": sorted({(page + 1) // 2 for page in document_pages}),
            "document_low_confidence": bool(profile.get("document_low_confidence")),
        })
    return {
        "cedula_candidate_count": len(roster_numbers | evidence_numbers),
        "cedula_candidate_pages": sorted(roster_by_page.keys() | evidence_by_page.keys()),
        "roster_candidate_count": len(roster_numbers),
        "roster_candidate_pages": sorted(roster_by_page),
        "evidence_candidate_count": len(evidence_numbers),
        "evidence_candidate_pages": sorted(evidence_by_page),
        "cedula_checksum_valid_count": len(roster_checksum_valid_numbers | evidence_checksum_valid_numbers),
        "cedula_checksum_invalid_count": len(roster_checksum_invalid_numbers | evidence_checksum_invalid_numbers),
        "roster_checksum_valid_count": len(roster_checksum_valid_numbers),
        "roster_checksum_invalid_count": len(roster_checksum_invalid_numbers),
        "evidence_checksum_valid_count": len(evidence_checksum_valid_numbers),
        "evidence_checksum_invalid_count": len(evidence_checksum_invalid_numbers),
        "candidate_matches": len(confident_matches),
        "roster_candidates_without_document_candidate": len(confident_roster_numbers - confident_evidence_numbers),
        "evidence_candidates_without_roster_candidate": len(confident_evidence_numbers - confident_roster_numbers),
        "low_confidence_evidence_candidates": len(low_confidence_evidence_candidates),
        "verified_identities": 0,
        "identity_association_requires_human_review": True,
        "participant_review_items": participant_review_items,
        "participants_requiring_review_count": sum(
            item["review_status"] == "REQUIERE_REVISION" for item in participant_review_items
        ),
        "participants_ok_automatic_count": sum(
            item["review_status"] == "OK_AUTOMATICO" for item in participant_review_items
        ),
    }


def sanitize_json(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(key): sanitize_json(val) for key, val in value.items()}
    if isinstance(value, list):
        return [sanitize_json(item) for item in value]
    if isinstance(value, tuple):
        return [sanitize_json(item) for item in value]
    if isinstance(value, set):
        return [sanitize_json(item) for item in value]
    return value


def resolve_relation_owner_ids(session, current_user: UserORM, owner_id: str | None) -> list[str]:
    if current_user.role == "admin":
        if owner_id:
            owner_exists = session.scalar(select(UserORM.id).where(UserORM.id == owner_id))
            if owner_exists is None:
                raise HTTPException(status_code=404, detail="Usuario no encontrado.")
            return [owner_id]
        return [user_id for user_id in session.scalars(select(UserORM.id)).all()]

    if owner_id and owner_id != current_user.id:
        raise HTTPException(status_code=403, detail="No tienes permiso para ver relaciones de otros usuarios.")
    return [current_user.id]


def build_relation_records(
    session,
    rows: list[tuple[ProcessingJobORM, HumanReviewORM | None]],
) -> list[Dict[str, Any]]:
    creator_ids = {job.created_by for job, _ in rows if job.created_by}
    creators = {
        user.id: user.username
        for user in session.scalars(select(UserORM).where(UserORM.id.in_(creator_ids))).all()
    } if creator_ids else {}

    action_codes = {
        (job.payload_metadata or {}).get("action_code")
        for job, _ in rows
        if (job.payload_metadata or {}).get("action_code")
    }
    action_names = {
        row.action_code: row.name
        for row in session.scalars(
            select(FormativeActionORM).where(FormativeActionORM.action_code.in_(action_codes))
        ).all()
    } if action_codes else {}

    relations: list[Dict[str, Any]] = []
    for job, review in rows:
        metadata = job.payload_metadata or {}
        action_code = metadata.get("action_code")
        relations.append({
            "job_id": job.id,
            "source_pdf_name": Path(job.source_pdf_path).name,
            "created_by": job.created_by,
            "created_by_username": creators.get(job.created_by or "", "—"),
            "status": job.status,
            "queued_at": job.queued_at.isoformat() if job.queued_at else None,
            "action_code": action_code,
            "action_name": action_names.get(action_code, "—"),
            "participant_count": metadata.get("roster_candidate_count") or metadata.get("cedula_candidate_count") or 0,
            "review": {
                "reviewer": review.reviewer,
                "comments": review.comments,
                "final_status": review.final_status,
                "reviewed_at": review.reviewed_at.isoformat() if review.reviewed_at else None,
            } if review else None,
        })
    return relations


def set_processing_progress(
    row: ProcessingJobORM,
    percent: int,
    stage: str,
    message: str,
) -> None:
    metadata = dict(row.payload_metadata or {})
    metadata["processing_progress_percent"] = max(0, min(100, int(percent)))
    metadata["processing_stage"] = stage
    metadata["processing_message"] = message
    if percent >= 100:
        metadata["processing_finished_at"] = datetime.now(timezone.utc).isoformat()
    elif "processing_finished_at" in metadata:
        metadata.pop("processing_finished_at", None)
    row.payload_metadata = sanitize_json(metadata)


def persist_processing_progress(
    job_id: str,
    percent: int,
    stage: str,
    message: str,
) -> None:
    with SessionLocal() as session:
        row = session.get(ProcessingJobORM, job_id)
        if row is None:
            return
        set_processing_progress(row, percent, stage, message)
        session.commit()


async def process_job_in_background(
    job_id: str,
    actor_username: str,
    force: bool = False,
    request_metadata: Dict[str, Any] | None = None,
) -> None:
    request_metadata = request_metadata or {}
    with SessionLocal() as session:
        row = session.get(ProcessingJobORM, job_id)
        if row is None:
            return
        row.status = JobStatus.PROCESSING.value
        row.started_at = datetime.now(timezone.utc)
        row.finished_at = None
        set_processing_progress(row, 5, "initializing", "Iniciando revisión automática.")
        session.commit()

        previous_review = session.get(HumanReviewORM, job_id)
        if previous_review is not None:
            reviewed_at = datetime.now(timezone.utc)
            session.add(AuditEventORM(
                id=str(uuid4()),
                job_id=job_id,
                actor=actor_username,
                action="HUMAN_REVIEW_INVALIDATED_BY_REPROCESS",
                details=sanitize_json({
                    "previous_reviewer": previous_review.reviewer,
                    "previous_final_status": previous_review.final_status,
                }),
                created_at=reviewed_at,
            ))
            session.delete(previous_review)
            session.commit()

        previous_verifications = session.scalars(
            select(ValidationResultORM).where(
                ValidationResultORM.job_id == job_id,
                ValidationResultORM.participant_id.is_not(None),
            )
        ).all()
        if previous_verifications:
            session.add(AuditEventORM(
                id=str(uuid4()),
                job_id=job_id,
                actor=actor_username,
                action="CEDULA_VERIFICATIONS_INVALIDATED_BY_REPROCESS",
                details={"invalidated_count": len(previous_verifications)},
                created_at=datetime.now(timezone.utc),
            ))
            for verification in previous_verifications:
                session.delete(verification)
            session.commit()

        try:
            set_processing_progress(row, 20, "reading_pdf", "Leyendo PDF y metadatos.")
            session.commit()

            pdf_metadata: Dict[str, Any] = {}
            pdf_text_source = "source_unavailable"
            text_extraction = None
            if row.source_pdf_path and is_managed_evidence_available(row.source_pdf_path):
                pdf_metadata = pdf_service.inspect_existing_pdf(row.source_pdf_path)
                persist_processing_progress(job_id, 45, "ocr", "Ejecutando OCR y extracción de texto.")
                ocr_last_saved_percent = 45

                def ocr_progress_callback(page_number: int, total_pages: int) -> None:
                    nonlocal ocr_last_saved_percent
                    if total_pages <= 0:
                        return
                    percent = 45 + int((page_number / total_pages) * 20)
                    if percent <= ocr_last_saved_percent and percent < 100:
                        return
                    ocr_last_saved_percent = percent
                    persist_processing_progress(
                        job_id,
                        percent,
                        "ocr",
                        f"Ejecutando OCR y extracción de texto ({page_number}/{total_pages}).",
                    )

                text_extraction = await run_in_threadpool(
                    pdf_service.extract_text,
                    row.source_pdf_path,
                    ocr_provider,
                    ocr_progress_callback,
                )
                pdf_text_source = text_extraction.source

            persist_processing_progress(job_id, 70, "analyzing", "Analizando candidatos y reglas.")
            candidate_summary = summarize_document_candidates(
                text_extraction.pages if text_extraction else []
            )
            participant_review_items = candidate_summary.get("participant_review_items", [])
            ai_review = None
            if ai_ocr_analyzer is not None and ai_ocr_analyzer.is_configured() and text_extraction is not None:
                candidate_pages = sorted(
                    set(
                        candidate_summary.get("roster_candidate_pages", [])
                        + candidate_summary.get("evidence_candidate_pages", [])
                    )
                )
                persist_processing_progress(job_id, 75, "ai_analysis", "Analizando resultados con IA asistida.")
                try:
                    ai_review = await run_in_threadpool(
                        ai_ocr_analyzer.analyze,
                        text_extraction,
                        focus_pages=candidate_pages or None,
                        participant_hints=participant_review_items,
                        action_code=pdf_metadata.get("action_code"),
                    )
                except Exception as exc:
                    candidate_summary["ai_ocr_enabled"] = False
                    candidate_summary["ai_ocr_error"] = str(exc)
                    candidate_summary["ai_ocr_warnings"] = [f"AI OCR no disponible: {exc}"]
                    persist_processing_progress(
                        job_id,
                        78,
                        "ai_analysis",
                        f"AI OCR no disponible para este documento: {exc}",
                    )
                else:
                    if ai_review is not None:
                        participant_review_items = ai_review.participant_items or participant_review_items
                        candidate_summary["ai_ocr_enabled"] = True
                        candidate_summary["ai_ocr_provider"] = ai_review.provider
                        candidate_summary["ai_ocr_model"] = ai_review.model
                        candidate_summary["ai_ocr_summary"] = ai_review.summary
                        candidate_summary["ai_ocr_confidence"] = ai_review.confidence
                        candidate_summary["ai_ocr_warnings"] = ai_review.warnings or []
                        candidate_summary["ai_participant_items"] = ai_review.participant_items
                        candidate_summary["participant_review_items"] = participant_review_items
                        candidate_summary["participants_requiring_review_count"] = sum(
                            item.get("review_status") == "REQUIERE_REVISION" for item in participant_review_items
                        )
                        candidate_summary["participants_ok_automatic_count"] = sum(
                            item.get("review_status") == "OK_AUTOMATICO" for item in participant_review_items
                        )
                    persist_processing_progress(job_id, 82, "ai_analysis", "IA asistida aplicada al documento.")
            else:
                candidate_summary["ai_ocr_enabled"] = False
                candidate_summary["ai_ocr_warnings"] = [ai_ocr_diagnostic] if ai_ocr_diagnostic else []

            row.payload_metadata = sanitize_json({
                **(row.payload_metadata or {}),
                **request_metadata,
                "force_reprocess": force,
                "page_count": pdf_metadata.get("page_count", (row.payload_metadata or {}).get("page_count")),
                "estimated_sheets_duplex": (
                    int((pdf_metadata.get("page_count") + 1) / 2)
                    if isinstance(pdf_metadata.get("page_count"), int)
                    else (row.payload_metadata or {}).get("estimated_sheets_duplex")
                ),
                "sha256": pdf_metadata.get("sha256", (row.payload_metadata or {}).get("sha256")),
                "action_code": pdf_metadata.get("action_code", (row.payload_metadata or {}).get("action_code")),
                "text_extraction_source": pdf_text_source,
                "ocr_provider": text_extraction.provider if text_extraction else None,
                "ocr_provider_diagnostic": ocr_provider_diagnostic if pdf_text_source == "ocr_not_configured" else None,
                "ocr_pages_processed": len(text_extraction.pages) if text_extraction else 0,
                "extracted_text_characters": len(text_extraction.full_text) if text_extraction else 0,
                "text_pages_detected": sum(bool(page.text.strip()) for page in text_extraction.pages) if text_extraction else 0,
                "ai_ocr_enabled": candidate_summary.get("ai_ocr_enabled", False),
                "ai_ocr_provider": candidate_summary.get("ai_ocr_provider"),
                "ai_ocr_model": candidate_summary.get("ai_ocr_model"),
                "ai_ocr_summary": candidate_summary.get("ai_ocr_summary"),
                "ai_ocr_confidence": candidate_summary.get("ai_ocr_confidence"),
                "ai_ocr_warnings": candidate_summary.get("ai_ocr_warnings", []),
                "ai_ocr_error": candidate_summary.get("ai_ocr_error"),
                "ai_participant_items": candidate_summary.get("ai_participant_items", []),
                **candidate_summary,
            })
            persist_processing_progress(job_id, 90, "saving", "Guardando resultados de revisión.")

            row.status = JobStatus.REQUIRES_REVIEW.value
            row.finished_at = datetime.now(timezone.utc)

            if pdf_text_source == "ocr_not_configured":
                message = (
                    "PDF conservado y analizado. Es un escaneo sin texto digital; "
                    "el motor OCR no está disponible en este servidor, por lo que no se validaron participantes ni documentos."
                )
                if ocr_provider_diagnostic:
                    message = f"{message} {ocr_provider_diagnostic}"
            elif pdf_text_source == "source_unavailable":
                message = "No se encontró el PDF original. El job requiere revisión y no se ejecutó la validación documental."
            elif pdf_text_source == "ocr":
                message = (
                    f"OCR local ({text_extraction.provider}) extrajo texto de "
                    f"{row.payload_metadata['text_pages_detected']} páginas. "
                    f"Detectó {candidate_summary['roster_candidate_count']} candidatos en páginas de lista "
                    f"y {candidate_summary['evidence_candidate_count']} en otras páginas. "
                    f"{candidate_summary['cedula_checksum_valid_count']} pasan el control matemático "
                    f"del dígito; {candidate_summary['cedula_checksum_invalid_count']} no lo pasan. "
                    f"Coincidencias exactas entre páginas con lectura confiable: {candidate_summary['candidate_matches']}; "
                    f"{candidate_summary['low_confidence_evidence_candidates']} candidatos de documento "
                    "tienen baja confianza de OCR. El control de dígito y las coincidencias no confirman "
                    "la identidad ni que el documento pertenezca a una persona; revisa los PDF originales."
                )
            else:
                message = (
                    f"Se extrajo texto digital de {row.payload_metadata['text_pages_detected']} páginas. "
                    f"Detectó {candidate_summary['roster_candidate_count']} candidatos en páginas de lista "
                    f"y {candidate_summary['evidence_candidate_count']} en otras páginas. "
                    f"{candidate_summary['cedula_checksum_valid_count']} pasan el control matemático "
                    f"del dígito; {candidate_summary['cedula_checksum_invalid_count']} no lo pasan. "
                    f"Coincidencias exactas entre páginas con lectura confiable: {candidate_summary['candidate_matches']}. "
                    "El control no confirma identidad ni asociación; revisa los PDF originales."
                )
            if pdf_metadata.get("action_code"):
                message = f"{message} Código candidato detectado en el nombre del archivo: {pdf_metadata.get('action_code')}."
            if row.payload_metadata.get("ai_ocr_enabled") and row.payload_metadata.get("ai_ocr_summary"):
                message = f"{message} IA asistida: {row.payload_metadata.get('ai_ocr_summary')}."
            total_pages = row.payload_metadata.get("page_count")
            processed_pages = row.payload_metadata.get("ocr_pages_processed")
            if isinstance(total_pages, int) and isinstance(processed_pages, int) and processed_pages > 0 and processed_pages < total_pages:
                message = (
                    f"{message} Modo rápido activo: OCR aplicado a {processed_pages} de {total_pages} páginas "
                    "(para acelerar revisión automática)."
                )

            summary_result = session.scalar(
                select(ValidationResultORM).where(
                    ValidationResultORM.job_id == job_id,
                    ValidationResultORM.participant_id.is_(None),
                    ValidationResultORM.identity_document_id.is_(None),
                    ValidationResultORM.rule_id.is_(None),
                ).limit(1)
            )
            if summary_result is None:
                summary_result = ValidationResultORM(
                    id=str(uuid4()),
                    job_id=job_id,
                    participant_id=None,
                    identity_document_id=None,
                    rule_id=None,
                    outcome=ResultOutcome.NO_VERIFICABLE.value,
                    reason_code=RuleCode.DIFERENCIA_DETECTADA.value,
                    message=message,
                    evidence_ref=Path(row.source_pdf_path).name,
                    created_at=datetime.now(timezone.utc),
                )
                session.add(summary_result)
            else:
                summary_result.outcome = ResultOutcome.NO_VERIFICABLE.value
                summary_result.reason_code = RuleCode.DIFERENCIA_DETECTADA.value
                summary_result.message = message
                summary_result.evidence_ref = Path(row.source_pdf_path).name
                summary_result.created_at = datetime.now(timezone.utc)

            set_processing_progress(row, 100, "done", "Revisión automática completada.")
            session.commit()
        except Exception as exc:
            row.status = JobStatus.FAILED.value
            row.finished_at = datetime.now(timezone.utc)
            set_processing_progress(row, 100, "failed", f"Error durante la revisión automática: {exc}")
            session.commit()


@app.get("/health")
async def health_check() -> Dict[str, str]:
    return {"status": "ok"}


@app.post("/api/v1/auth/login")
async def login(payload: LoginRequest, response: Response) -> Dict[str, str]:
    username = payload.username.strip().lower()
    with SessionLocal() as session:
        user = session.scalar(select(UserORM).where(UserORM.username == username))
        if user is None or not user.is_active or not verify_password(payload.password, user.password_hash):
            raise HTTPException(
                status_code=status.HTTP_401_UNAUTHORIZED,
                detail="Usuario o contraseña incorrectos.",
            )

        now = time.time()
        session.execute(delete(UserSessionORM).where(UserSessionORM.expires_at <= now))
        token = secrets.token_urlsafe(32)
        session.add(UserSessionORM(
            token_hash=hashlib.sha256(token.encode("utf-8")).hexdigest(),
            user_id=user.id,
            expires_at=now + AUTH_SESSION_SECONDS,
        ))
        session.commit()
        response.set_cookie(
            key=AUTH_COOKIE_NAME,
            value=token,
            httponly=True,
            secure=AUTH_COOKIE_SECURE,
            samesite="lax",
            max_age=AUTH_SESSION_SECONDS,
            path="/",
        )
        return {"username": user.username, "role": user.role}


@app.get("/api/v1/auth/me")
async def current_user_profile(current_user: UserORM = Depends(get_current_user)) -> Dict[str, str]:
    return {"username": current_user.username, "role": current_user.role}


@app.post("/api/v1/auth/logout", status_code=status.HTTP_204_NO_CONTENT)
async def logout(
    request: Request,
    current_user: UserORM = Depends(get_current_user),
) -> Response:
    del current_user
    token = request.cookies.get(AUTH_COOKIE_NAME)
    if token:
        token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
        with SessionLocal() as session:
            auth_session = session.get(UserSessionORM, token_hash)
            if auth_session is not None:
                session.delete(auth_session)
                session.commit()
    logout_response = Response(status_code=status.HTTP_204_NO_CONTENT)
    logout_response.delete_cookie(
        key=AUTH_COOKIE_NAME,
        path="/",
        secure=AUTH_COOKIE_SECURE,
        httponly=True,
        samesite="lax",
    )
    return logout_response


@app.post("/api/v1/jobs", status_code=status.HTTP_410_GONE)
async def create_job(
    payload: CreateJobRequest,
    current_user: UserORM = Depends(get_current_user),
) -> Dict[str, str]:
    del payload, current_user
    raise HTTPException(
        status_code=status.HTTP_410_GONE,
        detail="Creating jobs by server path is disabled. Upload the PDF through /api/v1/jobs/upload.",
    )


@app.post("/api/v1/jobs/upload", status_code=status.HTTP_201_CREATED)
async def upload_pdf_job(
    background_tasks: BackgroundTasks,
    file: UploadFile = File(...),
    metadata: str = Form(default='{}'),
    current_user: UserORM = Depends(get_current_user),
) -> Dict[str, Any]:
    if not file.filename or not file.filename.lower().endswith(".pdf"):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="Only PDF files are supported")

    if len(metadata.encode("utf-8")) > MAX_UPLOAD_METADATA_BYTES:
        raise HTTPException(status_code=status.HTTP_413_CONTENT_TOO_LARGE, detail="Upload metadata exceeds the 16 KB limit.")
    try:
        payload_metadata = json.loads(metadata) if metadata else {}
    except json.JSONDecodeError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="metadata must be valid JSON") from exc
    if not isinstance(payload_metadata, dict):
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail="metadata must be a JSON object")

    job_id = str(uuid4())
    try:
        stored = pdf_service.store_pdf(job_id, file.filename, file.file)
    except PdfUploadLimitError as exc:
        raise HTTPException(status_code=status.HTTP_413_CONTENT_TOO_LARGE, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(exc)) from exc
    payload_metadata = sanitize_json({
        **payload_metadata,
        "requested_by": current_user.username,
        "page_count": stored["page_count"],
        "action_code": stored["action_code"],
        "file_size": stored["file_size"],
        "processing_progress_percent": 0,
        "processing_stage": "queued",
        "processing_message": "PDF recibido. En cola para revisión automática.",
    })
    with SessionLocal() as session:
        job = ProcessingJobORM(
            id=job_id,
            source_pdf_path=stored["storage_path"],
            original_file_hash=str(stored["sha256"]),
            status=JobStatus.QUEUED.value,
            created_by=current_user.id,
            payload_metadata=payload_metadata,
            queued_at=datetime.now(timezone.utc),
        )
        session.add(job)
        session.commit()
    background_tasks.add_task(
        process_job_in_background,
        job_id,
        current_user.username,
        False,
        {},
    )
    with SessionLocal() as session:
        job = session.get(ProcessingJobORM, job_id)
        review = session.get(HumanReviewORM, job_id)
        return serialize_processing_job(job, review)


@app.get("/api/v1/jobs")
async def list_jobs(current_user: UserORM = Depends(get_current_user)) -> List[Dict[str, Any]]:
    with SessionLocal() as session:
        statement = (
            select(ProcessingJobORM, HumanReviewORM)
            .outerjoin(HumanReviewORM, HumanReviewORM.job_id == ProcessingJobORM.id)
            .order_by(ProcessingJobORM.queued_at.desc())
        )
        if current_user.role != "admin":
            statement = statement.where(ProcessingJobORM.created_by == current_user.id)
        rows = session.execute(statement).all()
        return [serialize_processing_job(job, review) for job, review in rows]


@app.get("/api/v1/dashboard/summary")
async def dashboard_summary(current_user: UserORM = Depends(get_current_user)) -> Dict[str, int]:
    with SessionLocal() as session:
        statement = (
            select(ProcessingJobORM, HumanReviewORM)
            .outerjoin(HumanReviewORM, HumanReviewORM.job_id == ProcessingJobORM.id)
        )
        if current_user.role != "admin":
            statement = statement.where(ProcessingJobORM.created_by == current_user.id)
        rows = session.execute(statement).all()
        available_jobs = [
            (job, review)
            for job, review in rows
            if is_managed_evidence_available(job.source_pdf_path)
        ]
        job_ids = [job.id for job, _ in available_jobs]
        results = session.scalars(
            select(ValidationResultORM).where(ValidationResultORM.job_id.in_(job_ids))
        ).all() if job_ids else []

        return {
            "total": len(available_jobs),
            "pending": sum(job.status in {JobStatus.QUEUED.value, JobStatus.PROCESSING.value} for job, _ in available_jobs),
            "in_review": sum(
                job.status == JobStatus.REQUIRES_REVIEW.value and review is None
                for job, review in available_jobs
            ),
            "incidents": sum(result.outcome == ResultOutcome.INCIDENCIA.value for result in results),
            "not_verifiable": sum(result.outcome == ResultOutcome.NO_VERIFICABLE.value for result in results),
            "processed": sum(job.status == JobStatus.COMPLETED.value for job, _ in available_jobs),
        }


@app.get("/api/v1/relations")
async def list_user_relations(
    owner_id: str | None = None,
    limit: int = 200,
    current_user: UserORM = Depends(get_current_user),
) -> Dict[str, Any]:
    limit = min(max(limit, 1), 500)
    with SessionLocal() as session:
        owner_ids = resolve_relation_owner_ids(session, current_user, owner_id)
        rows = session.execute(
            select(ProcessingJobORM, HumanReviewORM)
            .outerjoin(HumanReviewORM, HumanReviewORM.job_id == ProcessingJobORM.id)
            .where(ProcessingJobORM.created_by.in_(owner_ids))
            .order_by(ProcessingJobORM.queued_at.desc())
            .limit(limit)
        ).all()
        items = build_relation_records(session, rows)
        return {
            "items": items,
            "total": len(items),
            "owner_scope": "all" if current_user.role == "admin" and owner_id is None else "personal",
        }


@app.get("/api/v1/relations/returned")
async def list_returned_relations(
    owner_id: str | None = None,
    limit: int = 200,
    current_user: UserORM = Depends(get_current_user),
) -> Dict[str, Any]:
    limit = min(max(limit, 1), 500)
    with SessionLocal() as session:
        owner_ids = resolve_relation_owner_ids(session, current_user, owner_id)
        rows = session.execute(
            select(ProcessingJobORM, HumanReviewORM)
            .join(HumanReviewORM, HumanReviewORM.job_id == ProcessingJobORM.id)
            .where(
                ProcessingJobORM.created_by.in_(owner_ids),
                HumanReviewORM.final_status == "DEVUELTA",
            )
            .order_by(HumanReviewORM.reviewed_at.desc())
            .limit(limit)
        ).all()
        items = build_relation_records(session, rows)
        return {"items": items, "total": len(items)}


@app.get("/api/v1/relations/returned/{job_id}/print")
async def get_returned_relation_print_data(
    job_id: UUID,
    current_user: UserORM = Depends(get_current_user),
) -> Dict[str, Any]:
    with SessionLocal() as session:
        job = session.get(ProcessingJobORM, str(job_id))
        if job is None:
            raise HTTPException(status_code=404, detail="Expediente no encontrado.")
        if current_user.role != "admin" and job.created_by != current_user.id:
            raise HTTPException(status_code=404, detail="Expediente no encontrado.")

        review = session.get(HumanReviewORM, str(job_id))
        if review is None or review.final_status != "DEVUELTA":
            raise HTTPException(status_code=404, detail="Este expediente no está en estado DEVUELTA.")

        creator = session.get(UserORM, job.created_by) if job.created_by else None
        metadata = job.payload_metadata or {}
        action_code = metadata.get("action_code")
        if not isinstance(action_code, str) or not re.fullmatch(r"\d{4}-\d{5,8}", action_code):
            action_code = None
        action_name = "—"
        if action_code:
            action_row = session.scalar(select(FormativeActionORM).where(FormativeActionORM.action_code == action_code))
            if action_row is not None and action_row.name:
                action_name = action_row.name
        reviewed_at = review.reviewed_at.astimezone(timezone.utc) if review.reviewed_at else datetime.now(timezone.utc)
        participant_count = metadata.get("roster_candidate_count") or metadata.get("cedula_candidate_count") or 0

        return {
            "job_id": job.id,
            "codigo_accion_formativa": action_code or "—",
            "accion_formativa": action_name,
            "fecha_devuelto": reviewed_at.strftime("%d/%m/%y"),
            "hora_devuelto": reviewed_at.strftime("%H:%M"),
            "asesor_observacion": review.comments or "INICIO DEVUELTO, ACTUALIZAR",
            "participantes": participant_count,
            "devuelto_por": review.reviewer,
            "registrado_por_usuario": creator.username if creator else "—",
            "documento": Path(job.source_pdf_path).name,
        }


@app.get("/api/v1/formative-actions")
async def list_formative_actions(
    q: str = "",
    limit: int = 25,
    offset: int = 0,
    current_user: UserORM = Depends(get_current_user),
) -> Dict[str, Any]:
    del current_user
    limit = min(max(limit, 1), 100)
    offset = max(offset, 0)
    statement = select(FormativeActionORM)
    count_statement = select(func.count()).select_from(FormativeActionORM)
    query = q.strip()
    if query:
        pattern = f"%{query.lower()}%"
        condition = (
            func.lower(FormativeActionORM.action_code).like(pattern)
            | func.lower(FormativeActionORM.name).like(pattern)
        )
        statement = statement.where(condition)
        count_statement = count_statement.where(condition)

    with SessionLocal() as session:
        total = session.scalar(count_statement) or 0
        rows = session.scalars(
            statement.order_by(FormativeActionORM.action_code).offset(offset).limit(limit)
        ).all()
        latest_import = session.scalar(
            select(FormativeActionImportORM).order_by(FormativeActionImportORM.imported_at.desc()).limit(1)
        )
        return {
            "items": [{
                "action_code": row.action_code,
                "name": row.name,
                "start_date": row.start_date,
                "status": row.status,
                "fields": row.fields or {},
                "imported_at": row.imported_at.isoformat(),
                "imported_by": row.imported_by,
            } for row in rows],
            "total": total,
            "limit": limit,
            "offset": offset,
            "latest_import": {
                "filename": latest_import.filename,
                "imported_by": latest_import.imported_by,
                "imported_at": latest_import.imported_at.isoformat(),
                "total_rows": latest_import.total_rows,
                "inserted_rows": latest_import.inserted_rows,
                "updated_rows": latest_import.updated_rows,
                "columns": latest_import.columns or [],
            } if latest_import else None,
        }


@app.post("/api/v1/formative-actions/import")
async def import_formative_actions(
    file: UploadFile = File(...),
    current_user: UserORM = Depends(require_admin),
) -> Dict[str, Any]:
    filename = Path(file.filename or "").name
    content = await file.read(MAX_ACTION_IMPORT_BYTES + 1)
    if len(content) > MAX_ACTION_IMPORT_BYTES:
        raise HTTPException(status_code=413, detail="El archivo supera el límite de 20 MB.")
    records, columns = await run_in_threadpool(parse_action_import, filename, content)
    inserted = 0
    updated = 0
    imported_at = datetime.now(timezone.utc)
    with SessionLocal() as session:
        existing_by_code: dict[str, FormativeActionORM] = {}
        action_codes = [record["action_code"] for record in records]
        for start in range(0, len(action_codes), 500):
            existing_rows = session.scalars(
                select(FormativeActionORM).where(
                    FormativeActionORM.action_code.in_(action_codes[start:start + 500])
                )
            ).all()
            existing_by_code.update({row.action_code: row for row in existing_rows})

        for record in records:
            row = existing_by_code.get(record["action_code"])
            if row is None:
                row = FormativeActionORM(
                    action_code=record["action_code"],
                    name=record["name"],
                    start_date=record["start_date"],
                    status=record["status"],
                    fields=record["fields"],
                    imported_at=imported_at,
                    imported_by=current_user.username,
                )
                session.add(row)
                inserted += 1
            else:
                row.name = record["name"]
                row.start_date = record["start_date"]
                row.status = record["status"]
                row.fields = record["fields"]
                row.imported_at = imported_at
                row.imported_by = current_user.username
                updated += 1

        session.add(FormativeActionImportORM(
            id=str(uuid4()),
            filename=filename,
            file_hash=hashlib.sha256(content).hexdigest(),
            imported_by=current_user.username,
            imported_at=imported_at,
            total_rows=len(records),
            inserted_rows=inserted,
            updated_rows=updated,
            columns=columns,
        ))
        session.commit()
    return {
        "filename": filename,
        "total_rows": len(records),
        "inserted_rows": inserted,
        "updated_rows": updated,
        "imported_at": imported_at.isoformat(),
    }


@app.get("/api/v1/jobs/{job_id}")
async def get_job(job_id: UUID, current_user: UserORM = Depends(get_current_user)) -> Dict[str, Any]:
    with SessionLocal() as session:
        row = get_accessible_job(session, str(job_id), current_user)
        return serialize_processing_job(row, session.get(HumanReviewORM, str(job_id)))


@app.post("/api/v1/jobs/{job_id}/process")
async def process_job(
    job_id: UUID,
    payload: ProcessJobRequest,
    background_tasks: BackgroundTasks,
    current_user: UserORM = Depends(get_current_user),
) -> Dict[str, Any]:
    with SessionLocal() as session:
        row = get_accessible_job(session, str(job_id), current_user)
        if row.status == JobStatus.PROCESSING.value:
            progress = (row.payload_metadata or {}).get("processing_progress_percent", 0)
            stage = (row.payload_metadata or {}).get("processing_stage", "processing")
            message = (row.payload_metadata or {}).get("processing_message", "La revisión automática ya está en curso.")
            return {
                "job_id": str(job_id),
                "status": row.status,
                "progress_percent": progress,
                "stage": stage,
                "message": message,
            }
        row.status = JobStatus.PROCESSING.value
        row.started_at = datetime.now(timezone.utc)
        row.finished_at = None
        set_processing_progress(row, 5, "queued", "Revisión automática en cola.")
        session.commit()

    background_tasks.add_task(
        process_job_in_background,
        str(job_id),
        current_user.username,
        payload.force,
        payload.metadata,
    )
    return {
        "job_id": str(job_id),
        "status": JobStatus.PROCESSING.value,
        "progress_percent": 5,
        "stage": "queued",
        "message": "Revisión automática iniciada.",
    }


@app.get("/api/v1/jobs/{job_id}/results")
async def get_job_results(
    job_id: UUID,
    current_user: UserORM = Depends(get_current_user),
) -> List[Dict[str, Any]]:
    with SessionLocal() as session:
        get_accessible_job(session, str(job_id), current_user)
        results = session.scalars(select(ValidationResultORM).where(ValidationResultORM.job_id == str(job_id))).all()
        return [serialize_validation_result(result) for result in results]


@app.get("/api/v1/jobs/{job_id}/review")
async def get_job_review(
    job_id: UUID,
    current_user: UserORM = Depends(get_current_user),
) -> Dict[str, Any] | None:
    with SessionLocal() as session:
        get_accessible_job(session, str(job_id), current_user)
        review = session.get(HumanReviewORM, str(job_id))
        if review is None:
            return None
        return {
            "job_id": review.job_id,
            "reviewer": review.reviewer,
            "comments": review.comments,
            "final_status": review.final_status,
            "reviewed_at": review.reviewed_at.isoformat(),
        }


@app.put("/api/v1/jobs/{job_id}/review")
async def save_job_review(
    job_id: UUID,
    payload: HumanReviewRequest,
    current_user: UserORM = Depends(get_current_user),
) -> Dict[str, Any]:
    reviewer = current_user.username
    if payload.final_status is not None and payload.final_status not in FINAL_STATUSES:
        raise HTTPException(
            status_code=status.HTTP_422_UNPROCESSABLE_CONTENT,
            detail=f"final_status must be one of: {', '.join(sorted(FINAL_STATUSES))}",
        )

    with SessionLocal() as session:
        get_accessible_job(session, str(job_id), current_user)

        reviewed_at = datetime.now(timezone.utc)
        review = session.get(HumanReviewORM, str(job_id))
        if review is None:
            review = HumanReviewORM(job_id=str(job_id), reviewer=reviewer)
            session.add(review)

        review.reviewer = reviewer
        review.comments = payload.comments.strip()
        review.final_status = payload.final_status
        review.reviewed_at = reviewed_at
        audit_event = AuditEventORM(
            id=str(uuid4()),
            job_id=str(job_id),
            actor=reviewer,
            action="HUMAN_REVIEW_SAVED",
            details=sanitize_json({
                "comments": review.comments,
                "final_status": review.final_status,
            }),
            created_at=reviewed_at,
        )
        session.add(audit_event)
        session.commit()
        return {
            "job_id": review.job_id,
            "reviewer": review.reviewer,
            "comments": review.comments,
            "final_status": review.final_status,
            "reviewed_at": review.reviewed_at.isoformat(),
        }


@app.get("/api/v1/jobs/{job_id}/evidence")
async def get_job_evidence(
    job_id: UUID,
    current_user: UserORM = Depends(get_current_user),
) -> FileResponse:
    with SessionLocal() as session:
        job = get_accessible_job(session, str(job_id), current_user)
        evidence_path = job.source_pdf_path

    path = Path(evidence_path).resolve()
    try:
        path.relative_to(pdf_service.storage_dir.resolve())
    except ValueError as exc:
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Evidence is outside the managed PDF storage") from exc
    if not path.is_file():
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Original PDF evidence not found")
    audit_evidence_access(
        str(job_id),
        current_user.username,
        "EVIDENCE_PDF_REQUESTED",
        {"resource": "original_pdf"},
    )
    return FileResponse(
        path,
        media_type="application/pdf",
        filename=path.name,
        content_disposition_type="inline",
        headers={
            "Cache-Control": "private, no-store",
            "X-Content-Type-Options": "nosniff",
        },
    )


@app.get("/api/v1/jobs/{job_id}/preview")
async def get_job_evidence_preview(
    job_id: UUID,
    page: int = 1,
    current_user: UserORM = Depends(get_current_user),
) -> Response:
    if page < 1:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_CONTENT, detail="Page must be 1 or greater.")
    with SessionLocal() as session:
        job = get_accessible_job(session, str(job_id), current_user)
        evidence_path = Path(job.source_pdf_path).resolve()

    try:
        evidence_path.relative_to(pdf_service.storage_dir.resolve())
    except ValueError as exc:
        raise HTTPException(
            status_code=status.HTTP_403_FORBIDDEN,
            detail="Evidence is outside the managed PDF storage",
        ) from exc
    if not evidence_path.is_file():
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Original PDF evidence not found")

    image, page_count = await run_in_threadpool(render_pdf_page_preview, str(evidence_path), page)
    audit_evidence_access(
        str(job_id),
        current_user.username,
        "EVIDENCE_PREVIEW_GENERATED",
        {"page": page},
    )
    return Response(
        content=image,
        media_type="image/png",
        headers={
            "Cache-Control": "no-store",
            "X-Page-Count": str(page_count),
        },
    )


@app.get("/api/v1/jobs/{job_id}/audit")
async def get_job_audit(
    job_id: UUID,
    current_user: UserORM = Depends(get_current_user),
) -> List[Dict[str, Any]]:
    with SessionLocal() as session:
        get_accessible_job(session, str(job_id), current_user)
        events = session.scalars(
            select(AuditEventORM)
            .where(AuditEventORM.job_id == str(job_id))
            .order_by(AuditEventORM.created_at.desc())
        ).all()
        return [{
            "id": event.id,
            "job_id": event.job_id,
            "actor": event.actor,
            "action": event.action,
            "details": event.details or {},
            "created_at": event.created_at.isoformat(),
        } for event in events]


@app.delete("/api/v1/jobs/{job_id}")
async def delete_job(
    job_id: UUID,
    current_user: UserORM = Depends(get_current_user),
) -> Dict[str, str]:
    if current_user.role != "admin":
        raise HTTPException(status_code=status.HTTP_403_FORBIDDEN, detail="Administrator role required")
    with SessionLocal() as session:
        row = session.get(ProcessingJobORM, str(job_id))
        if row is None:
            raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Processing job not found")
        session.delete(row)
        session.query(ValidationResultORM).filter(ValidationResultORM.job_id == str(job_id)).delete()
        session.commit()
        return {"status": "deleted", "job_id": str(job_id)}
